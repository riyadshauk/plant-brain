"""Repeatable import of the supplied weekly PDFs. Never writes to source PDFs."""
from __future__ import annotations

import io
import json
import re
import unicodedata
from difflib import SequenceMatcher
from collections import defaultdict
from pathlib import Path

import pymupdf
import openpyxl
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "generated_data"
IMAGE_DIR = ROOT / "app" / "public" / "course-images"
COLLECTION_ID = "plant-materials-2026"
MODULE_NAMES = {
    1: "Herbaceous Perennials", 2: "Grasses and Palms", 3: "Flowering Shrubs",
    4: "Trees", 5: "Groundcovers and Vines", 6: "Agaves, Aloes, Cacti and Succulents",
    7: "Special Interest and Character", 8: "Hedges and Screens", 9: "California Natives",
}
EXCLUDED_IMAGE_IDS = {"w5-p3-i1"}  # Black placeholder embedded on the Carmel Sur slide.


def compact(value: str | None) -> str:
    return " ".join((value or "").replace("\n", " ").split())


def slug(value: str) -> str:
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "-", value).strip("-")


def identity(name: str) -> str:
    """Conservative identity. Only the documented Coffeeberry synonym is merged."""
    name = compact(name).lower()
    if name in {"rhamnus californica", "rhamnus (frangula) californica"}:
        return "rhamnus-californica"
    return slug(name)


def stable_id(week: int, position: int, name: str, id_map: dict, known_names: dict[str, str]) -> str:
    """Keep prior IDs across reorderings and small source spelling corrections."""
    if name in known_names:
        return known_names[name]
    old = id_map.get(f"week-{week}-position-{position}")
    if old and SequenceMatcher(None, old["name"].casefold(), name.casefold()).ratio() >= 0.85:
        return old["id"]
    return identity(name)


def size_parts(raw: str) -> tuple[str | None, str | None]:
    text = compact(raw)
    if not text:
        return None, None
    match = re.match(r"^(.+?)\s*[xX×]\s*(.+)$", text)
    if not match:
        return (text, None) if text.lower() not in {"varies", "spreading"} else (None, None)
    return compact(match.group(1)), compact(match.group(2))


def slide_row_index(week: int, page: int) -> int | None:
    i = page - 2
    if week == 2 and page in (24, 25):
        return 23 if page == 24 else 22
    if week == 3 and page >= 15:
        return i + 1
    if week == 4 and page in (16, 17):
        return 15 if page == 16 else 14
    if week == 8:
        if page in (14, 15):
            return 11
        if page >= 16:
            return page - 4 if page <= 17 else page - 3
    if week == 9:
        if page == 2:
            return None  # A. glauca is narrower than the generic list entry.
        return i if page <= 7 else i + 1  # No Ceanothus slide.
    return i


def main() -> None:
    OUTPUT.mkdir(exist_ok=True)
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    plants: dict[str, dict] = {}
    memberships: list[dict] = []
    images: list[dict] = []
    unmatched_slides: list[dict] = []
    rows_by_week: dict[int, list[dict]] = {}
    issues: list[str] = []
    id_map_path = OUTPUT / "plant-id-map.json"
    id_map = json.loads(id_map_path.read_text(encoding="utf-8")) if id_map_path.exists() else {}
    known_names = {record["name"]: record["id"] for record in id_map.values()}
    next_id_map = {}
    master_book = openpyxl.load_workbook(ROOT / "Master Plant List.xlsx", data_only=True)
    master_rows = {}
    for excel_row, cells in enumerate(master_book.active.iter_rows(values_only=True), start=1):
        if isinstance(cells[0], (int, float)) and 1 <= cells[0] <= 9 and cells[1]:
            key = (int(cells[0]), compact(str(cells[1])))
            master_rows[key] = (excel_row, [compact(str(x)) if x is not None else "" for x in cells[1:6]])

    for week in MODULE_NAMES:
        source = next(ROOT.glob(f"Week {week} Plant List*.pdf"))
        doc = pymupdf.open(source)
        tables = doc[0].find_tables().tables
        if len(tables) != 1:
            raise ValueError(f"Expected one table: {source.name}")
        rows = []
        for row in tables[0].extract():
            if not row or compact(row[0]) != str(week):
                continue
            weekly_values = [compact(cell) for cell in row[1:6]]
            key = (week, weekly_values[0])
            if key not in master_rows:
                raise ValueError(f"Weekly entry absent from master list: {key}")
            excel_row, master_values = master_rows[key]
            latin, common, dimension, origin, factor = master_values
            for field, master_value, weekly_value in zip(
                ("scientific name", "common name", "dimensions", "origin", "plant factor"),
                master_values, weekly_values,
            ):
                if master_value.casefold() != weekly_value.casefold():
                    issues.append(f"Week {week} {latin}: {field} differs — master {master_value!r}; weekly PDF {weekly_value!r}. Imported master value.")
            if not latin or not common:
                raise ValueError(f"Missing name: {source.name}: {row}")
            if factor not in {"H", "M", "L", "VL"}:
                raise ValueError(f"Invalid plant factor: {source.name}: {row}")
            height, spread = size_parts(dimension)
            plant_id = stable_id(week, len(rows) + 1, latin, id_map, known_names)
            next_id_map[f"week-{week}-position-{len(rows)+1}"] = {"name": latin, "id": plant_id}
            known_names[latin] = plant_id
            if plant_id not in plants:
                plants[plant_id] = {
                    "id": plant_id, "scientificName": latin, "commonName": common,
                    "alternateCourseNames": [], "imageIds": [], "sourceRefs": [],
                }
            elif latin != plants[plant_id]["scientificName"]:
                plants[plant_id]["alternateCourseNames"].append(latin)
            source_ref = {"file": "Master Plant List.xlsx", "sheet": master_book.active.title, "row": excel_row}
            plants[plant_id]["sourceRefs"].append(source_ref)
            category = MODULE_NAMES[week]
            if week == 2:
                category = "Grasses, sedges and rushes" if len(rows) < 12 else "Palms"
            if week == 5:
                category = "Groundcovers" if len(rows) < 11 else "Vines"
            membership = {
                "id": f"{COLLECTION_ID}-w{week}-{len(rows)+1}",
                "collectionId": COLLECTION_ID, "moduleId": f"week-{week}",
                "plantId": plant_id, "order": len(rows) + 1,
                "category": category,
                "facts": {
                    "scientificName": latin, "commonName": common,
                    "dimensionsRaw": dimension, "height": height, "spread": spread,
                    "origin": origin, "wucolsZone3": factor,
                    "distinguishingFeatures": [], "importantFacts": [],
                    **({"flowerColor": compact(row[6])} if len(row) > 6 and compact(row[6]) else {}),
                },
                "sourceRef": source_ref,
                "supportingSourceRef": {"file": source.name, "page": 1},
            }
            memberships.append(membership)
            rows.append(membership)
        rows_by_week[week] = rows

    if len(memberships) != len(master_rows):
        raise ValueError(f"Master/weekly entry mismatch: {len(master_rows)} vs {len(memberships)}")

    for week in MODULE_NAMES:
        source = next(ROOT.glob(f"*Week {week}*Presentation.pdf"))
        doc = pymupdf.open(source)
        rows = rows_by_week[week]
        for page_index in range(1, len(doc)):
            page_num = page_index + 1
            page = doc[page_index]
            title = compact(page.get_text())
            row_index = slide_row_index(week, page_num)
            membership = rows[row_index] if row_index is not None and row_index < len(rows) else None
            if membership is None:
                unmatched_slides.append({"week": week, "file": source.name, "page": page_num, "title": title})
            else:
                membership.setdefault("presentationRefs", []).append({"file": source.name, "page": page_num, "title": title})
            for number, img in enumerate(page.get_images(full=True), start=1):
                image_id = f"w{week}-p{page_num}-i{number}"
                filename = f"{image_id}.webp"
                if image_id in EXCLUDED_IMAGE_IDS:
                    (IMAGE_DIR / filename).unlink(missing_ok=True)
                    continue
                raw = doc.extract_image(img[0])["image"]
                im = Image.open(io.BytesIO(raw)).convert("RGB")
                im.save(IMAGE_DIR / filename, "WEBP", quality=87, method=4)
                image = {
                    "id": image_id, "plantId": membership["plantId"] if membership else None,
                    "url": f"/course-images/{filename}", "part": None,
                    "source": "instructor presentation", "caption": title,
                    "notes": None, "rights": "Instructor course material; personal study use",
                    "sourceRef": {"file": source.name, "page": page_num, "imageNumber": number},
                }
                images.append(image)
                if membership:
                    plants[membership["plantId"]]["imageIds"].append(image_id)

    for plant in plants.values():
        plant["imageIds"] = list(dict.fromkeys(plant["imageIds"]))
        plant["sourceRefs"] = list({(r["file"], r["row"]): r for r in plant["sourceRefs"]}.values())

    for week, rows in rows_by_week.items():
        unpictured = [r["facts"]["scientificName"] for r in rows if not r.get("presentationRefs")]
        if unpictured:
            issues.append(f"Week {week} list entries without a matched presentation slide: {', '.join(unpictured)}.")
    if unmatched_slides:
        issues.append("Unassociated presentation slides: " + "; ".join(
            f"Week {s['week']} p.{s['page']} ({s['title']})" for s in unmatched_slides
        ) + ".")

    course = {
        "schemaVersion": 1,
        "collections": [{
            "id": COLLECTION_ID, "name": "Plant Materials · 2026",
            "description": "Instructor-supplied weekly plant lists and presentations.",
            "moduleIds": [f"week-{n}" for n in MODULE_NAMES],
            "studyProfile": {
                "requiredFacts": ["scientificName", "commonName", "height", "spread", "wucolsZone3", "distinguishingFeatures", "importantFacts"],
                "exam": {"plantCount": 15, "fields": ["scientificName", "commonName", "height", "spread", "wucolsZone3", "distinguishingFeatures", "importantFacts"]},
            },
        }],
        "modules": [{"id": f"week-{n}", "collectionId": COLLECTION_ID, "name": f"Week {n}", "topic": name, "order": n} for n, name in MODULE_NAMES.items()],
        "plants": list(plants.values()), "memberships": memberships,
        "images": images, "unmatchedSlides": unmatched_slides,
    }
    course_json = json.dumps(course, indent=2, ensure_ascii=False) + "\n"
    (OUTPUT / "course.json").write_text(course_json, encoding="utf-8")
    (ROOT / "app" / "public" / "course.json").write_text(course_json, encoding="utf-8")
    id_map_path.write_text(json.dumps(next_id_map, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    stats = {
        "weeklyEntries": len(memberships), "uniquePlants": len(plants), "modules": len(rows_by_week),
        "extractedImages": len(images), "associatedImages": sum(x["plantId"] is not None for x in images),
        "plantsWithoutImages": sum(not p["imageIds"] for p in plants.values()),
    }
    (OUTPUT / "import-report.json").write_text(json.dumps({"stats": stats, "issues": issues, "unmatchedSlides": unmatched_slides}, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(stats, indent=2))
    for issue in issues:
        print(issue)


if __name__ == "__main__":
    main()
